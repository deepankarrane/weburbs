from io import BytesIO

import pandas as pd
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.views.decorators.http import require_GET
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import Color

from projects.api.helper import get_project
from projects.api.simulate import with_t0
from projects.models import (
    BuySellPrice,
    Commodity,
    Demand,
    DSM,
    ProcComDir,
    Process,
    ProcessCommodity,
    Site,
    Storage,
    SupIm,
    TimeVarEff,
    Transmission,
)

GLOBAL_DESCRIPTIONS = {
    "CO2 limit": (
        "Limits the sum of all created (as calculated by commodity_balance) "
        "CO2 in all sites; Only relevant if not minimized"
    ),
    "Cost limit": (
        "Limits the sum of all costs in all sites; Only relevant if not minimized"
    ),
}

SITE_COLUMNS = ["Name", "area"]
COMMODITY_COLUMNS = ["Site", "Commodity", "Type", "price", "max", "maxperhour"]
PROCESS_COLUMNS = [
    "Site",
    "Process",
    "inst-cap",
    "cap-lo",
    "cap-up",
    "max-grad",
    "min-fraction",
    "inv-cost",
    "fix-cost",
    "var-cost",
    "wacc",
    "depreciation",
    "area-per-cap",
]
PROCESS_COMMODITY_COLUMNS = ["Process", "Commodity", "Direction", "ratio", "ratio-min"]
STORAGE_COLUMNS = [
    "Site",
    "Storage",
    "Commodity",
    "inst-cap-c",
    "cap-lo-c",
    "cap-up-c",
    "inst-cap-p",
    "cap-lo-p",
    "cap-up-p",
    "eff-in",
    "eff-out",
    "inv-cost-p",
    "inv-cost-c",
    "fix-cost-p",
    "fix-cost-c",
    "var-cost-p",
    "var-cost-c",
    "wacc",
    "depreciation",
    "init",
    "discharge",
    "ep-ratio",
]
TRANSMISSION_COLUMNS = [
    "Site In",
    "Site Out",
    "Transmission",
    "Commodity",
    "eff",
    "inv-cost",
    "fix-cost",
    "var-cost",
    "inst-cap",
    "cap-lo",
    "cap-up",
    "wacc",
    "depreciation",
    "reactance",
    "difflimit",
    "base_voltage",
]
DSM_COLUMNS = ["Site", "Commodity", "delay", "eff", "recov", "cap-max-do", "cap-max-up"]


# Sheets where blank cells stay blank instead of being written as #N/A
NO_NA_SHEETS = {"Global"}

# URBS input template color scheme (theme colors from single_year_example.xlsx)
FILL_HEADER_PARAM = PatternFill(
    patternType="solid",
    fgColor=Color(theme=5, tint=0.3999755851924192),
)
FILL_HEADER_TIMESERIES = PatternFill(
    patternType="solid",
    fgColor=Color(theme=6, tint=0.5999938962981048),
)
FILL_HEADER_GLOBAL_LABEL = PatternFill(
    patternType="solid",
    fgColor=Color(theme=0, tint=0.0),
)
FILL_DATA_IDENTIFIER = PatternFill(
    patternType="solid",
    fgColor=Color(theme=6, tint=0.5999938962981048),
)
FILL_DATA_PARAM = PatternFill(
    patternType="solid",
    fgColor=Color(theme=4, tint=0.7999816888943144),
)
FILL_DATA_TIMESTEP = PatternFill(
    patternType="solid",
    fgColor=Color(theme=9, tint=0.3999755851924192),
)

SHEET_ID_COLUMN_COUNT = {
    "Site": 1,
    "Commodity": 3,
    "Process": 2,
    "Process-Commodity": 3,
    "Transmission": 4,
    "Storage": 3,
    "DSM": 2,
}
TIMESERIES_SHEETS = {"Demand", "SupIm", "Buy-Sell-Price", "TimeVarEff"}


def _apply_cell_fill(cell, fill):
    if fill is not None:
        cell.fill = fill


def apply_global_sheet_style(ws):
    for col in range(1, ws.max_column + 1):
        if col == 1:
            _apply_cell_fill(ws.cell(1, col), FILL_HEADER_GLOBAL_LABEL)
        else:
            _apply_cell_fill(ws.cell(1, col), FILL_HEADER_PARAM)

    for row in range(2, ws.max_row + 1):
        _apply_cell_fill(ws.cell(row, 1), FILL_DATA_IDENTIFIER)
        for col in range(2, ws.max_column + 1):
            _apply_cell_fill(ws.cell(row, col), FILL_DATA_PARAM)


def apply_table_sheet_style(ws, id_column_count):
    for col in range(1, ws.max_column + 1):
        if col <= id_column_count:
            _apply_cell_fill(ws.cell(1, col), None)
        else:
            _apply_cell_fill(ws.cell(1, col), FILL_HEADER_PARAM)

    for row in range(2, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            if col <= id_column_count:
                _apply_cell_fill(ws.cell(row, col), FILL_DATA_IDENTIFIER)
            else:
                _apply_cell_fill(ws.cell(row, col), FILL_DATA_PARAM)


def apply_timeseries_sheet_style(ws):
    _apply_cell_fill(ws.cell(1, 1), None)
    for col in range(2, ws.max_column + 1):
        _apply_cell_fill(ws.cell(1, col), FILL_HEADER_TIMESERIES)

    for row in range(2, ws.max_row + 1):
        _apply_cell_fill(ws.cell(row, 1), FILL_DATA_TIMESTEP)
        for col in range(2, ws.max_column + 1):
            _apply_cell_fill(ws.cell(row, col), FILL_DATA_PARAM)


def apply_urbs_sheet_style(ws, sheet_name):
    if ws.max_row == 0 or ws.max_column == 0:
        return

    if sheet_name == "Global":
        apply_global_sheet_style(ws)
    elif sheet_name in TIMESERIES_SHEETS:
        apply_timeseries_sheet_style(ws)
    elif sheet_name in SHEET_ID_COLUMN_COUNT:
        apply_table_sheet_style(ws, SHEET_ID_COLUMN_COUNT[sheet_name])


def write_sheet(df, writer, sheet_name):
    df = df.replace(-1, "inf")
    if sheet_name not in NO_NA_SHEETS:
        df = df.fillna("#N/A")
    df.to_excel(writer, sheet_name=sheet_name, index=False)
    apply_urbs_sheet_style(writer.sheets[sheet_name], sheet_name)


def build_timeseries_frame(series_by_name):
    max_steps = 0
    for values in series_by_name.values():
        if len(values) > max_steps:
            max_steps = len(values)

    frame = {"t": list(range(max_steps))}
    for name, values in series_by_name.items():
        padded = list(values) + [None] * (max_steps - len(values))
        frame[name] = padded

    return pd.DataFrame(frame)


@login_required
@require_GET
def download(request, project_name):
    project = get_project(request.user, project_name)

    sites = Site.objects.filter(project=project).order_by("name")
    commodities = Commodity.objects.filter(site__project=project).select_related("site").order_by(
        "site__name", "name"
    )
    processes = Process.objects.filter(site__project=project).select_related("site").order_by(
        "site__name", "name"
    )
    process_commodities = (
        ProcessCommodity.objects.filter(process__site__project=project)
        .select_related("process", "commodity")
        .order_by("process__site__name", "process__name", "commodity__name")
    )
    storages = Storage.objects.filter(site__project=project).select_related(
        "site", "commodity"
    ).order_by("site__name", "name")
    transmissions = Transmission.objects.filter(commodityin__site__project=project).select_related(
        "commodityin__site", "commodityout__site"
    ).order_by("commodityin__site__name", "commodityout__site__name", "commodityin__name")
    dsms = DSM.objects.filter(commodity__site__project=project).select_related(
        "commodity__site"
    ).order_by("commodity__site__name", "commodity__name")
    demands = Demand.objects.filter(commodity__site__project=project).select_related(
        "commodity__site"
    ).order_by("commodity__site__name", "commodity__name", "id")
    supims = SupIm.objects.filter(commodity__site__project=project).select_related(
        "commodity__site"
    ).order_by("commodity__site__name", "commodity__name", "id")
    buysell_prices = BuySellPrice.objects.filter(project=project).order_by("name", "id")
    tve_items = TimeVarEff.objects.filter(process__site__project=project).select_related(
        "process__site"
    ).order_by("process__site__name", "process__name", "id")

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(
            [
                {
                    "Property": "CO2 limit",
                    "value": project.co2limit,
                    "description": GLOBAL_DESCRIPTIONS["CO2 limit"],
                },
                {
                    "Property": "Cost limit",
                    "value": project.costlimit,
                    "description": GLOBAL_DESCRIPTIONS["Cost limit"],
                },
            ],
            columns=["Property", "value", "description"],
        ).pipe(write_sheet, writer, "Global")

        pd.DataFrame(
            [{"Name": site.name, "area": site.area} for site in sites],
            columns=SITE_COLUMNS,
        ).pipe(write_sheet, writer, "Site")

        pd.DataFrame(
            [
                {
                    "Site": commodity.site.name,
                    "Commodity": commodity.name,
                    "Type": commodity.get_com_type_label(),
                    "price": commodity.price,
                    "max": commodity.max,
                    "maxperhour": commodity.maxperhour,
                }
                for commodity in commodities
            ],
            columns=COMMODITY_COLUMNS,
        ).pipe(write_sheet, writer, "Commodity")

        pd.DataFrame(
            [
                {
                    "Site": process.site.name,
                    "Process": process.name,
                    "inst-cap": process.instcap,
                    "cap-lo": process.caplo,
                    "cap-up": process.capup,
                    "max-grad": process.maxgrad,
                    "min-fraction": process.minfraction,
                    "inv-cost": process.invcost,
                    "fix-cost": process.fixcost,
                    "var-cost": process.varcost,
                    "wacc": process.wacc,
                    "depreciation": process.depreciation,
                    "area-per-cap": process.areapercap,
                }
                for process in processes
            ],
            columns=PROCESS_COLUMNS,
        ).pipe(write_sheet, writer, "Process")

        pd.DataFrame(
            [
                {
                    "Process": process_commodity.process.name,
                    "Commodity": process_commodity.commodity.name,
                    "Direction": ProcComDir(process_commodity.direction).name,
                    "ratio": process_commodity.ratio,
                    "ratio-min": process_commodity.ratiomin,
                }
                for process_commodity in process_commodities
            ],
            columns=PROCESS_COMMODITY_COLUMNS,
        ).pipe(write_sheet, writer, "Process-Commodity")

        pd.DataFrame(
            [
                {
                    "Site": storage.site.name,
                    "Storage": storage.name,
                    "Commodity": storage.commodity.name,
                    "inst-cap-c": storage.instcapc,
                    "cap-lo-c": storage.caploc,
                    "cap-up-c": storage.capupc,
                    "inst-cap-p": storage.instcapp,
                    "cap-lo-p": storage.caplop,
                    "cap-up-p": storage.capupp,
                    "eff-in": storage.effin,
                    "eff-out": storage.effout,
                    "inv-cost-p": storage.invcostp,
                    "inv-cost-c": storage.invcostc,
                    "fix-cost-p": storage.fixcostp,
                    "fix-cost-c": storage.fixcostc,
                    "var-cost-p": storage.varcostp,
                    "var-cost-c": storage.varcostc,
                    "wacc": storage.wacc,
                    "depreciation": storage.depreciation,
                    "init": storage.init,
                    "discharge": storage.discharge,
                    "ep-ratio": storage.epratio,
                }
                for storage in storages
            ],
            columns=STORAGE_COLUMNS,
        ).pipe(write_sheet, writer, "Storage")

        demand_series = {}
        for demand in demands:
            key = f"{demand.commodity.site.name}.{demand.commodity.name}"
            if key not in demand_series:
                demand_series[key] = with_t0(demand.steps)
        build_timeseries_frame(demand_series).pipe(write_sheet, writer, "Demand")

        supim_series = {}
        for supim in supims:
            key = f"{supim.commodity.site.name}.{supim.commodity.name}"
            if key not in supim_series:
                supim_series[key] = with_t0(supim.steps)
        build_timeseries_frame(supim_series).pipe(write_sheet, writer, "SupIm")

        pd.DataFrame(
            [
                {
                    "Site In": transmission.commodityin.site.name,
                    "Site Out": transmission.commodityout.site.name,
                    "Transmission": transmission.get_trans_type_label(),
                    "Commodity": transmission.commodityin.name,
                    "eff": transmission.eff,
                    "inv-cost": transmission.invcost,
                    "fix-cost": transmission.fixcost,
                    "var-cost": transmission.varcost,
                    "inst-cap": transmission.instcap,
                    "cap-lo": transmission.caplo,
                    "cap-up": transmission.capup,
                    "wacc": transmission.wacc,
                    "depreciation": transmission.depreciation,
                    "reactance": transmission.reactance,
                    "difflimit": transmission.difflimit,
                    "base_voltage": transmission.basevoltage,
                }
                for transmission in transmissions
            ],
            columns=TRANSMISSION_COLUMNS,
        ).pipe(write_sheet, writer, "Transmission")

        pd.DataFrame(
            [
                {
                    "Site": dsm.commodity.site.name,
                    "Commodity": dsm.commodity.name,
                    "delay": dsm.delay,
                    "eff": dsm.eff,
                    "recov": dsm.recov,
                    "cap-max-do": dsm.capmaxdo,
                    "cap-max-up": dsm.capmaxup,
                }
                for dsm in dsms
            ],
            columns=DSM_COLUMNS,
        ).pipe(write_sheet, writer, "DSM")

        bsp_series = {}
        for bsp in buysell_prices:
            if bsp.name not in bsp_series:
                bsp_series[bsp.name] = with_t0(bsp.steps)
        build_timeseries_frame(bsp_series).pipe(write_sheet, writer, "Buy-Sell-Price")

        tve_series = {}
        for tve in tve_items:
            key = f"{tve.process.site.name}.{tve.process.name}"
            if key not in tve_series:
                tve_series[key] = with_t0(tve.steps)
        build_timeseries_frame(tve_series).pipe(write_sheet, writer, "TimeVarEff")

    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{project.name}.xlsx"'
    return response
